"""
Build the standard English & Arabic templates shipped with the app
(apps/api/templates/standard-contract.docx and standard-annex.docx).
Companies start from these or upload their own under Admin → Templates.

    pip install python-docx && python3 apps/api/scripts/make-standard-templates.py
"""
from pathlib import Path
import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent.parent / 'templates'
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
GREY = 'EEF2F6'


def rtl(p):
    ppr = p._p.get_or_add_pPr()
    ppr.append(OxmlElement('w:bidi'))
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in p.runs:
        r._r.get_or_add_rPr().append(OxmlElement('w:rtl'))


def shade(cell, fill=GREY):
    tcpr = cell._tc.get_or_add_tcPr()
    el = OxmlElement('w:shd')
    el.set(qn('w:val'), 'clear'); el.set(qn('w:color'), 'auto'); el.set(qn('w:fill'), fill)
    tcpr.append(el)


def write(cell, text, *, bold=False, arabic=False, size=10, align=None):
    p = cell.paragraphs[0]
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if arabic:
        rtl(p)
    elif align:
        p.alignment = align
    return p


def fix_widths(t, widths):
    t.autofit = False
    for i, w in enumerate(widths):
        t.columns[i].width = Cm(w)
        for cell in t.columns[i].cells:
            cell.width = Cm(w)


def new_doc():
    d = docx.Document()
    s = d.sections[0]
    s.page_height, s.page_width = Cm(29.7), Cm(21)
    s.left_margin = s.right_margin = Cm(1.8)
    s.top_margin, s.bottom_margin = Cm(1.5), Cm(1.5)
    st = d.styles['Normal']
    st.font.name = 'Arial'; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn('w:cs'), 'Arial')
    st.paragraph_format.space_after = Pt(2)
    footer = s.footer.paragraphs[0]
    footer.text = '{company_name} · {doc_no}'
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)
    return d


def heading(d, en, ar):
    t = d.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    write(t.cell(0, 0), '{company_name}', bold=True, size=14)
    p = write(t.cell(0, 1), ar, bold=True, size=16, arabic=True)
    t.cell(0, 0).add_paragraph(en).runs[0].font.size = Pt(16)
    t.cell(0, 0).paragraphs[1].runs[0].bold = True
    t.cell(0, 0).paragraphs[1].runs[0].font.color.rgb = ACCENT
    p2 = t.cell(0, 1).add_paragraph('{company_name}')
    rtl(p2)
    info = d.add_table(rows=1, cols=2)
    write(info.cell(0, 0), 'No. {doc_no}    Date {issue_date}')
    write(info.cell(0, 1), 'رقم {doc_no}    التاريخ {issue_date}', arabic=True)
    d.add_paragraph()


def facts(d, rows):
    t = d.add_table(rows=0, cols=3)
    t.style = 'Table Grid'
    for en, value, ar in rows:
        c = t.add_row().cells
        write(c[0], en, bold=True); shade(c[0])
        write(c[1], value, align=WD_ALIGN_PARAGRAPH.CENTER)
        write(c[2], ar, bold=True, arabic=True); shade(c[2])
    fix_widths(t, [4.2, 9, 4.2])
    d.add_paragraph()


def loop_table(d, headers, cells, widths):
    t = d.add_table(rows=2, cols=len(headers))
    t.style = 'Table Grid'
    for i, (en, ar) in enumerate(headers):
        c = t.cell(0, i); shade(c)
        write(c, en, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        p = c.add_paragraph(ar); rtl(p); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].bold = True
    for i, text in enumerate(cells):
        arabic = text.startswith('{service}')
        write(t.cell(1, i), text, arabic=arabic, align=WD_ALIGN_PARAGRAPH.CENTER if '{no}' in text or '{qty}' in text else None)
    fix_widths(t, widths)
    d.add_paragraph()


def two_lang(d, en, ar, *, bold=False):
    t = d.add_table(rows=1, cols=2)
    write(t.cell(0, 0), en, bold=bold)
    write(t.cell(0, 1), ar, bold=bold, arabic=True)


def signatures(d):
    d.add_paragraph()
    t = d.add_table(rows=3, cols=2)
    write(t.cell(0, 0), 'For {company_name} / الطرف الأول', bold=True)
    write(t.cell(0, 1), 'The client / الطرف الثاني', bold=True)
    write(t.cell(1, 0), 'Name / الاسم: {issued_by}')
    write(t.cell(1, 1), 'Name / الاسم: {client_name}')
    write(t.cell(2, 0), '\nSignature / التوقيع: ______________________')
    write(t.cell(2, 1), '\nSignature / التوقيع: ______________________')


EVENT_FACTS = [
    ('Client', '{client_name}', 'الطرف الثاني'),
    ('Phone', '{client_phone}', 'رقم الهاتف'),
    ('Civil ID', '{client_civil_id}', 'رقم البطاقة المدنية'),
    ('Address', '{client_address}', 'العنوان'),
    ('Event date', '{event_date}  {event_day_en} / {event_day}', 'تاريخ الحفل'),
    ('Venue', '{venue}', 'مكان الحفل'),
    ('Guests', '{pax}', 'عدد الأشخاص'),
    ('Event type', '{event_type}', 'نوع المناسبة'),
]

# ---- contract / proposal ----------------------------------------------------
d = new_doc()
heading(d, '{doc_title_en}', '{doc_title}')
facts(d, [('Contract ID', '{contract_no}', 'رقم العقد')] + EVENT_FACTS)
two_lang(d, 'Agreed services', 'الخدمات المتفق عليها', bold=True)
loop_table(d,
           [('#', '#'), ('Service', 'نوع العمل'), ('Details', 'التفاصيل'), ('Qty', 'العدد'), ('Price', 'السعر'), ('Amount', 'المبلغ')],
           ['{#lines}{no}', '{work_type}', '{details}', '{qty}', '{price}', '{amount}{/lines}'],
           [1, 4.2, 5.4, 1.6, 2.2, 2.6])
facts(d, [
    ('Total', '{total} {currency}', 'الإجمالي'),
    ('Down payment', '{down_payment} {currency}', 'المقدم / العربون'),
    ('Remaining', '{remaining} {currency}', 'الباقي'),
])
two_lang(d, 'Amount in words: {total_words}', 'المبلغ كتابة: {total_words_ar}')
d.add_paragraph()
two_lang(d, 'Notes: {notes}', 'ملاحظات: {notes}')
d.add_paragraph()
two_lang(d, 'Terms', 'الشروط', bold=True)
for en, ar in [
    ('The down payment is due on signature and is not refundable.', 'يُدفع المقدم عند التوقيع ولا يُرد.'),
    ('The remaining amount is due before the event date.', 'يُسدد المبلغ المتبقي قبل تاريخ الحفل.'),
    ('Any change to the services above is agreed in writing.', 'أي تعديل على الخدمات أعلاه يتم الاتفاق عليه كتابيًا.'),
    ('Venue fees charged by the hall or hotel are not included unless listed above.', 'رسوم القاعة أو الفندق غير مشمولة ما لم تُذكر أعلاه.'),
]:
    two_lang(d, f'• {en}', f'• {ar}')
signatures(d)
d.save(OUT / 'standard-contract.docx')

# ---- annex ------------------------------------------------------------------
a = new_doc()
heading(a, '{doc_title_en}', '{doc_title}')
facts(a, [('Contract ID', '{contract_no}', 'رقم العقد')] + EVENT_FACTS)
loop_table(a,
           [('#', '#'), ('Service', 'نوع العمل'), ('Service (EN)', 'الخدمة'), ('Qty', 'العدد'), ('Notes', 'التفاصيل')],
           ['{#annex}{no}', '{service}', '{service_alt}', '{qty}', '{notes}{/annex}'],
           [1, 4.6, 4.2, 1.6, 6])
two_lang(a, 'The client approves all the details above.', 'يوافق الطرف الثاني على جميع التفاصيل المذكورة أعلاه.', bold=True)
signatures(a)
a.save(OUT / 'standard-annex.docx')
print('written', OUT)
